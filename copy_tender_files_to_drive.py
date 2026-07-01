from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
REPORTS = ROOT / "reports"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Copy tender request files referenced by an Excel report to a target drive."
    )
    parser.add_argument(
        "--excel",
        type=Path,
        help="Path to the Excel report. Defaults to the newest reports/*Заявки*.xlsx file.",
    )
    parser.add_argument(
        "--dest",
        type=Path,
        default=Path(r"Z:\Тендеры"),
        help=r"Destination folder, for example Z:\Тендеры.",
    )
    parser.add_argument(
        "--mode",
        choices=["docs", "full"],
        default="docs",
        help="docs copies procurement docs and key linked files; full copies each analysis folder.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Show what would be copied without writing files.",
    )
    parser.add_argument(
        "--json",
        action="store_true",
        help="Print full manifest JSON to stdout.",
    )
    return parser.parse_args()


def newest_report() -> Path:
    candidates = [
        path
        for path in REPORTS.glob("*.xlsx")
        if not path.name.startswith("~$") and "Заявки" in path.name
    ]
    if not candidates:
        raise FileNotFoundError("No reports/*Заявки*.xlsx files found.")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def safe_name(value: object, max_len: int = 110) -> str:
    text = str(value or "").strip()
    text = re.sub(r"[<>:\"/\\|?*\x00-\x1f]+", "_", text)
    text = re.sub(r"\s+", " ", text).strip(" ._")
    return (text or "request")[:max_len].rstrip(" ._")


def compact_run_name(excel: Path) -> str:
    match = re.search(
        r"(\d{4})-(\d{2})-(\d{2})_(\d{4})-(\d{2})-(\d{2})",
        excel.stem,
    )
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if match:
        y1, m1, d1, y2, m2, d2 = match.groups()
        return f"{y1}{m1}{d1}-{y2}{m2}{d2}_{timestamp}"
    return f"{safe_name(excel.stem, 30)}_{timestamp}"


def cell_target(sheet, row: int, column: int) -> str:
    if not column:
        return ""
    cell = sheet.cell(row=row, column=column)
    if cell.hyperlink and cell.hyperlink.target:
        return str(cell.hyperlink.target)
    return str(cell.value or "")


def existing_path(value: str) -> Path | None:
    if not value or value.startswith(("http://", "https://")):
        return None
    value = value.strip().strip('"')
    if value.startswith("file:///"):
        value = value[8:]
    if re.match(r"^/[A-Za-z]:/", value):
        value = value[1:]
    value = value.replace("/", "\\") if re.match(r"^[A-Za-z]:/", value) else value
    path = Path(value)
    if path.exists():
        return path
    return None


def copy_file(src: Path, dst_dir: Path, copied: list[dict], dry_run: bool) -> Path | None:
    if not src.exists() or not src.is_file():
        return None
    dst_dir.mkdir(parents=True, exist_ok=True) if not dry_run else None
    dst = unique_path(dst_dir / src.name)
    if not dry_run:
        shutil.copy2(src, dst)
    copied.append({"source": str(src), "destination": str(dst), "type": "file"})
    return dst


def copy_dir(src: Path, dst: Path, copied: list[dict], dry_run: bool) -> Path | None:
    if not src.exists() or not src.is_dir():
        return None
    dst = unique_path(dst)
    if dry_run:
        count = sum(1 for path in src.rglob("*") if path.is_file())
    else:
        shutil.copytree(src, dst)
        count = sum(1 for path in dst.rglob("*") if path.is_file())
    copied.append({"source": str(src), "destination": str(dst), "type": "dir", "files": count})
    return dst


def unique_path(path: Path) -> Path:
    if not path.exists():
        return path
    stem = path.stem
    suffix = path.suffix
    parent = path.parent
    for index in range(2, 1000):
        candidate = parent / f"{stem}_{index}{suffix}"
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"Could not create unique path for {path}")


def header_columns(sheet) -> dict[str, int]:
    result: dict[str, int] = {}
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(row=1, column=column).value
        if value:
            result[str(value).strip()] = column
    return result


def first_column(headers: dict[str, int], *names: str) -> int:
    for name in names:
        if name in headers:
            return headers[name]
    return 0


def row_urls(sheet, row: int, headers: dict[str, int]) -> list[str]:
    urls: list[str] = []
    url_columns = [
        first_column(headers, "Ссылки в интернете"),
        first_column(headers, "Ссылка 1"),
        first_column(headers, "Ссылка 2"),
        first_column(headers, "Ссылка 3"),
    ]
    for column in [column for column in url_columns if column]:
        value = str(sheet.cell(row=row, column=column).value or "")
        urls.extend(line.strip() for line in value.splitlines() if line.strip().startswith(("http://", "https://")))
        cell = sheet.cell(row=row, column=column)
        if cell.hyperlink and cell.hyperlink.target.startswith(("http://", "https://")):
            urls.append(cell.hyperlink.target)
    result = []
    seen = set()
    for url in urls:
        if url not in seen:
            seen.add(url)
            result.append(url)
    return result


def set_report_link(sheet, row: int, column: int, target: Path | None, label: str | None = None) -> None:
    if not column or not target:
        return
    cell = sheet.cell(row=row, column=column)
    if label:
        cell.value = label
    cell.hyperlink = str(target)
    cell.style = "Hyperlink"


def rewrite_report_links(report_path: Path, updates: list[dict]) -> None:
    workbook = load_workbook(report_path)
    sheet = workbook.active
    for update in updates:
        row = update["row"]
        for column, target, label in update["links"]:
            set_report_link(sheet, row, column, Path(target), label)
    workbook.save(report_path)


def same_path(left: Path | None, right: Path | None) -> bool:
    if not left or not right:
        return False
    try:
        return str(left.resolve()).lower() == str(right.resolve()).lower()
    except OSError:
        return str(left).lower() == str(right).lower()


def main() -> int:
    args = parse_args()
    excel = (args.excel or newest_report()).resolve()
    dest_root = args.dest.resolve()
    if not excel.exists():
        raise FileNotFoundError(excel)
    if not dest_root.drive:
        raise ValueError(f"Destination must include a drive/path: {dest_root}")
    if not args.dry_run:
        dest_root.mkdir(parents=True, exist_ok=True)

    run_root = dest_root / compact_run_name(excel)
    report_copy = run_root / excel.name
    if not args.dry_run:
        run_root.mkdir(parents=True, exist_ok=True)
        shutil.copy2(excel, report_copy)

    workbook = load_workbook(excel, data_only=False)
    sheet = workbook.active
    headers = header_columns(sheet)
    columns = {
        "email_text": first_column(headers, "Письмо"),
        "tz": first_column(headers, "ТЗ"),
        "note": first_column(headers, "Аналитическая записка"),
        "estimate": first_column(headers, "Смета себестоимости"),
        "docs": first_column(headers, "Документация"),
        "request_report": first_column(headers, "Excel-отчет запроса"),
        "result_dir": first_column(headers, "Папка результата (путь)", "Папка результата"),
    }
    manifest = {
        "source_excel": str(excel),
        "destination": str(run_root),
        "mode": args.mode,
        "dry_run": args.dry_run,
        "rows": [],
    }
    report_updates: list[dict] = []

    for row in range(2, sheet.max_row + 1):
        number = sheet.cell(row=row, column=1).value or row - 1
        object_name = sheet.cell(row=row, column=2).value or ""
        row_dir = run_root / f"{int(number):02d}_{safe_name(object_name, 28)}"
        copied: list[dict] = []
        urls = row_urls(sheet, row, headers)

        if not args.dry_run:
            row_dir.mkdir(parents=True, exist_ok=True)
            (row_dir / "links.txt").write_text("\n".join(urls), encoding="utf-8")

        result_dir = existing_path(cell_target(sheet, row, columns["result_dir"]))
        docs_dir = existing_path(cell_target(sheet, row, columns["docs"]))
        tz_path = existing_path(cell_target(sheet, row, columns["tz"]))
        note_path = existing_path(cell_target(sheet, row, columns["note"]))
        estimate_path = existing_path(cell_target(sheet, row, columns["estimate"]))
        email_text = existing_path(cell_target(sheet, row, columns["email_text"]))
        request_report = existing_path(cell_target(sheet, row, columns["request_report"]))
        row_links: list[tuple[int, str, str | None]] = []
        replacement_targets: list[tuple[Path, Path, str]] = []

        if args.mode == "full" and result_dir:
            copied_path = copy_dir(result_dir, row_dir / "analysis", copied, args.dry_run)
            if copied_path:
                row_links.append((columns["result_dir"], str(copied_path), "папка"))
                replacement_targets.append((result_dir, copied_path, "папка"))
        else:
            if result_dir:
                row_links.append((columns["result_dir"], str(row_dir), "папка"))
                replacement_targets.append((result_dir, row_dir, "папка"))
            if docs_dir:
                copied_path = copy_dir(docs_dir, row_dir / "documentation", copied, args.dry_run)
                if copied_path:
                    row_links.append((columns["docs"], str(copied_path), "документы"))
                    replacement_targets.append((docs_dir, copied_path, "документы"))
            elif result_dir and (result_dir / "01_downloaded_docs").exists():
                copied_path = copy_dir(result_dir / "01_downloaded_docs", row_dir / "documentation", copied, args.dry_run)
                if copied_path:
                    row_links.append((columns["docs"], str(copied_path), "документы"))
            copied_path = copy_file(tz_path, row_dir / "tz", copied, args.dry_run) if tz_path else None
            if copied_path:
                row_links.append((columns["tz"], str(copied_path), "ТЗ"))
                replacement_targets.append((tz_path, copied_path, "ТЗ"))
            copied_path = copy_file(email_text, row_dir, copied, args.dry_run) if email_text else None
            if copied_path:
                row_links.append((columns["email_text"], str(copied_path), "текст письма"))
                replacement_targets.append((email_text, copied_path, "текст письма"))
            copied_path = copy_file(request_report, row_dir, copied, args.dry_run) if request_report else None
            if copied_path:
                row_links.append((columns["request_report"], str(copied_path), "отчет"))
                replacement_targets.append((request_report, copied_path, "отчет"))

        copied_path = copy_file(note_path, row_dir / "analytics", copied, args.dry_run) if note_path else None
        if copied_path:
            row_links.append((columns["note"], str(copied_path), "записка"))
            replacement_targets.append((note_path, copied_path, "записка"))
        copied_path = copy_file(estimate_path, row_dir / "estimate", copied, args.dry_run) if estimate_path else None
        if copied_path:
            row_links.append((columns["estimate"], str(copied_path), "смета"))
            replacement_targets.append((estimate_path, copied_path, "смета"))
        for column in range(1, sheet.max_column + 1):
            local_target = existing_path(cell_target(sheet, row, column))
            if not local_target:
                continue
            for source_path, copied_path, label in replacement_targets:
                if same_path(local_target, source_path):
                    row_links.append((column, str(copied_path), label))
                    break
        if row_links:
            report_updates.append({"row": row, "links": row_links})

        manifest["rows"].append(
            {
                "number": number,
                "object": object_name,
                "customer": sheet.cell(row=row, column=3).value or "",
                "deadline": sheet.cell(row=row, column=9).value or "",
                "destination": str(row_dir),
                "urls": urls,
                "copied": copied,
            }
        )

    manifest_path = run_root / "manifest.json"
    if not args.dry_run:
        rewrite_report_links(report_copy, report_updates)
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    if args.json:
        print(json.dumps(manifest, ensure_ascii=False, indent=2))
    else:
        copied_items = sum(len(row["copied"]) for row in manifest["rows"])
        copied_dirs = sum(1 for row in manifest["rows"] for item in row["copied"] if item["type"] == "dir")
        copied_files = sum(1 for row in manifest["rows"] for item in row["copied"] if item["type"] == "file")
        print(f"Excel: {excel}")
        print(f"Destination: {run_root}")
        print(f"Rows: {len(manifest['rows'])}")
        print(f"Planned/copied items: {copied_items} ({copied_dirs} folders, {copied_files} files)")
        if args.dry_run:
            print("Dry run: no files were copied.")
        else:
            print(f"Manifest: {manifest_path}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
