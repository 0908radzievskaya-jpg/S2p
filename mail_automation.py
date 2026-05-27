from __future__ import annotations

import argparse
import base64
import datetime as dt
import email
import hashlib
import imaplib
import json
import mimetypes
import os
import re
import socket
import sys
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass, field
from email.header import decode_header, make_header
from email.message import Message
from html.parser import HTMLParser
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
from xml.etree import ElementTree as ET

try:
    from openpyxl import Workbook, load_workbook
except Exception as exc:  # pragma: no cover
    raise SystemExit("openpyxl is required to build the daily Excel report") from exc

try:
    import requests
except Exception:  # pragma: no cover
    requests = None

try:
    from bs4 import BeautifulSoup
except Exception:  # pragma: no cover
    BeautifulSoup = None

PDF_READER = None
for module_name in ("pypdf", "PyPDF2"):
    try:  # pragma: no cover
        PDF_READER = __import__(module_name, fromlist=["PdfReader"]).PdfReader
        break
    except Exception:
        continue


SECTION_KEYWORDS = {
    "здания": ["здание", "корпус", "административн", "жилой дом", "офис"],
    "сооружения": ["сооружение", "эстакад", "резервуар", "мост", "башня"],
    "дороги": ["дорог", "проезд", "трасс", "улиц", "развязк"],
    "площадки": ["площадк", "участок", "производственная база"],
    "благоустройство": ["благоустрой", "озеленен", "малые архитектурные формы"],
    "генеральный план": ["генплан", "генеральный план", "посадка", "схема планировочной организации"],
    "архитектурные решения": ["архитектур", "ар", "фасад", "планировочн"],
    "конструктивные решения": ["конструктив", "кр", "кж", "км", "основани", "фундамент"],
    "технология": ["технолог", "тх", "оборудован", "производственн"],
    "наружные инженерные сети": ["наружные сети", "нвк", "нв", "теплотрас", "наружное электроснабжение"],
    "внутренние инженерные сети": ["внутренние сети", "инженерные системы", "виc", "вк", "ов"],
    "электрика": ["электрик", "электроснабж", "эом", "эс", "силовое оборудование"],
    "связь": ["связь", "сс", "скс", "видеонаблюден", "пожарная сигнализация"],
    "отопление": ["отоплен", "ио", "итп"],
    "вентиляция": ["вентиляц", "кондиционирован", "ов"],
    "водоснабжение": ["водоснабжен", "хвс", "гвс"],
    "канализация": ["канализац", "водоотведен", "к1", "к2"],
    "тепловые сети": ["тепловые сети", "тс", "теплоснабжен"],
    "газоснабжение": ["газоснабжен", "газ", "грп", "грс"],
    "ПОС": ["пос", "проект организации строительства", "стройгенплан"],
    "сметные и коммерческие материалы": ["смет", "коммерческ", "кп", "тендер", "бюджет", "стоимость"],
}

FIELD_PATTERNS = {
    "address": [
        r"(?:адрес|местоположение|расположение)\s*[:\-]\s*(.+)",
    ],
    "stage": [
        r"(?:стадия|этап)\s*[:\-]\s*(.+)",
    ],
    "deadline": [
        r"(?:срок|сроки|дедлайн|deadline)\s*[:\-]\s*(.+)",
    ],
    "customer": [
        r"(?:заказчик)\s*[:\-]\s*(.+)",
    ],
    "work_type": [
        r"(?:вид работ|работы|предмет)\s*[:\-]\s*(.+)",
    ],
}

REPORT_HEADERS = [
    "Дата обработки",
    "Почтовый ящик",
    "UID письма",
    "Тема",
    "Объект",
    "Заказчик/отправитель",
    "Контакты",
    "Адрес/местоположение",
    "Вид работ",
    "Стадия",
    "Сроки",
    "Состав документации",
    "Классификация разделов",
    "Исходные файлы",
    "Ссылки",
    "Ключевые технические параметры",
    "Объемы",
    "Требования",
    "Ограничения",
    "Недостающие сведения",
    "Риски",
    "Рекомендация следующего действия",
    "Статус достаточности",
    "Папка материалов",
]

RUNTIME_CONFIG: Dict[str, object] = {"imap_timeout_seconds": 20}


class LinkExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.links: List[str] = []

    def handle_starttag(self, tag: str, attrs: List[Tuple[str, Optional[str]]]) -> None:
        if tag.lower() != "a":
            return
        for key, value in attrs:
            if key.lower() == "href" and value:
                self.links.append(value)


@dataclass
class MailboxConfig:
    name: str
    host: str
    port: int
    username: str
    password: str
    folder: str = "INBOX"
    use_ssl: bool = True


@dataclass
class ProcessedMessage:
    mailbox: str
    uid: str
    message_id: str
    subject: str
    sender: str
    date_folder: str
    output_dir: str
    attachments: List[str] = field(default_factory=list)
    links: List[str] = field(default_factory=list)
    downloaded_files: List[str] = field(default_factory=list)
    sections: List[str] = field(default_factory=list)
    extracted_fields: Dict[str, str] = field(default_factory=dict)
    notes: List[str] = field(default_factory=list)


@dataclass
class DownloadedResource:
    source_url: str
    final_url: str
    local_path: str
    content_type: str = ""
    text_content: str = ""


def slugify(value: str, fallback: str = "unknown") -> str:
    normalized = re.sub(r"\s+", "_", value.strip().lower())
    normalized = re.sub(r"[^0-9a-zA-Zа-яА-ЯёЁ._-]+", "_", normalized)
    normalized = normalized.strip("._-")
    return normalized or fallback


def decode_mime(value: Optional[str]) -> str:
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return value


def ensure_dir(path: Path) -> None:
    path.mkdir(parents=True, exist_ok=True)


def load_config(path: Path) -> Dict[str, object]:
    with path.open("r", encoding="utf-8") as fh:
        data = json.load(fh)
    data["mailboxes"] = [MailboxConfig(**mailbox) for mailbox in data.get("mailboxes", [])]
    return data


def load_state(state_path: Path) -> Dict[str, Dict[str, object]]:
    if not state_path.exists():
        return {"processed": {}}
    with state_path.open("r", encoding="utf-8") as fh:
        return json.load(fh)


def save_state(state_path: Path, state: Dict[str, Dict[str, object]]) -> None:
    ensure_dir(state_path.parent)
    with state_path.open("w", encoding="utf-8") as fh:
        json.dump(state, fh, ensure_ascii=False, indent=2)


def connect_mailbox(cfg: MailboxConfig) -> imaplib.IMAP4:
    timeout = int(RUNTIME_CONFIG.get("imap_timeout_seconds", 20))
    client = (
        imaplib.IMAP4_SSL(cfg.host, cfg.port, timeout=timeout)
        if cfg.use_ssl
        else imaplib.IMAP4(cfg.host, cfg.port, timeout=timeout)
    )
    client.login(cfg.username, cfg.password)
    status, _ = client.select(cfg.folder, readonly=True)
    if status != "OK":
        raise RuntimeError(f"Unable to select folder {cfg.folder!r} for mailbox {cfg.name}")
    return client


def search_all_uids(client: imaplib.IMAP4) -> List[str]:
    status, data = client.uid("search", None, "ALL")
    if status != "OK" or not data or not data[0]:
        return []
    return data[0].decode("utf-8", errors="ignore").split()


def fetch_message_bytes(client: imaplib.IMAP4, uid: str) -> bytes:
    status, data = client.uid("fetch", uid, "(RFC822)")
    if status != "OK" or not data:
        raise RuntimeError(f"Unable to fetch message UID {uid}")
    for item in data:
        if isinstance(item, tuple) and len(item) >= 2:
            return item[1]
    raise RuntimeError(f"Empty fetch response for UID {uid}")


def parse_email_date(message: Message) -> dt.date:
    raw = message.get("Date")
    if not raw:
        return dt.datetime.now().date()
    try:
        parsed = email.utils.parsedate_to_datetime(raw)
        if parsed.tzinfo:
            parsed = parsed.astimezone()
        return parsed.date()
    except Exception:
        return dt.datetime.now().date()


def extract_message_bodies(message: Message) -> Tuple[str, str]:
    text_parts: List[str] = []
    html_parts: List[str] = []
    if message.is_multipart():
        for part in message.walk():
            disposition = (part.get("Content-Disposition") or "").lower()
            if "attachment" in disposition:
                continue
            content_type = part.get_content_type()
            payload = part.get_payload(decode=True) or b""
            charset = part.get_content_charset() or "utf-8"
            try:
                decoded = payload.decode(charset, errors="ignore")
            except Exception:
                decoded = payload.decode("utf-8", errors="ignore")
            if content_type == "text/plain":
                text_parts.append(decoded)
            elif content_type == "text/html":
                html_parts.append(decoded)
    else:
        payload = message.get_payload(decode=True) or b""
        charset = message.get_content_charset() or "utf-8"
        decoded = payload.decode(charset, errors="ignore")
        if message.get_content_type() == "text/html":
            html_parts.append(decoded)
        else:
            text_parts.append(decoded)
    return "\n".join(text_parts), "\n".join(html_parts)


def html_to_text(html: str) -> str:
    if not html:
        return ""
    if BeautifulSoup is not None:
        return BeautifulSoup(html, "html.parser").get_text("\n", strip=True)
    return re.sub(r"<[^>]+>", " ", html)


DOCUMENT_EXTENSIONS = {
    ".pdf", ".doc", ".docx", ".xls", ".xlsx", ".xlsm", ".csv", ".zip", ".rar", ".7z",
    ".dwg", ".dxf", ".ifc", ".rvt", ".txt", ".rtf", ".xml", ".json", ".ppt", ".pptx",
}

IGNORED_ASSET_EXTENSIONS = {
    ".js", ".css", ".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp", ".ico", ".woff",
    ".woff2", ".ttf", ".map", ".mp4", ".webm", ".mp3",
}


def extract_links(text: str, html: str, base_url: str = "") -> List[str]:
    links = set(re.findall(r"https?://[^\s<>()\"']+", text or "", flags=re.IGNORECASE))
    parser = LinkExtractor()
    try:
        parser.feed(html or "")
    except Exception:
        pass
    for link in parser.links:
        resolved = urllib.parse.urljoin(base_url, link) if base_url else link
        if resolved.startswith("http://") or resolved.startswith("https://"):
            links.add(resolved)
    normalized_links = {normalize_candidate_url(link) for link in links}
    return sorted(link for link in normalized_links if link and is_useful_material_link(link))


def scope_host(hostname: str) -> str:
    host = (hostname or "").strip().lower()
    if not host:
        return ""
    if re.fullmatch(r"\d+\.\d+\.\d+\.\d+", host):
        return host
    parts = host.split(".")
    if len(parts) >= 2:
        return ".".join(parts[-2:])
    return host


def pad_base64(value: str) -> str:
    padding = len(value) % 4
    if padding:
        value += "=" * (4 - padding)
    return value


def strip_control_chars(value: str) -> str:
    return re.sub(r"[\x00-\x1F\x7F]", "", value or "").strip()


def try_decode_redirect_payload(value: str) -> Optional[str]:
    if not value:
        return None
    for encoding in ("utf-8", "cp1251"):
        try:
            decoded = strip_control_chars(
                base64.urlsafe_b64decode(pad_base64(value)).decode(encoding, errors="ignore")
            )
        except Exception:
            continue
        if decoded.startswith("http://") or decoded.startswith("https://"):
            return decoded
    return None


def normalize_candidate_url(url: str) -> str:
    if not url:
        return url
    url = strip_control_chars(url)
    parsed = urllib.parse.urlparse(url)
    query = urllib.parse.parse_qs(parsed.query)
    hostname = (parsed.hostname or "").lower()

    for key in ("u", "url", "target", "redirect", "to"):
        raw_value = (query.get(key) or [""])[0]
        unquoted = strip_control_chars(urllib.parse.unquote(raw_value))
        if unquoted.startswith("http://") or unquoted.startswith("https://"):
            return normalize_candidate_url(unquoted)
        candidate = try_decode_redirect_payload(raw_value)
        if candidate:
            return normalize_candidate_url(candidate)

    if hostname == "stat.techmail.pik.ru":
        for key in ("h", "u"):
            candidate = try_decode_redirect_payload((query.get(key) or [""])[0])
            if candidate:
                return normalize_candidate_url(candidate)
    cleaned_query: List[Tuple[str, str]] = []
    for key, values in query.items():
        if key.lower().startswith("utm_") or key.lower() in {"u_ik", "yclid", "ymclid"}:
            continue
        for value in values:
            cleaned_query.append((key, value))
    parsed = parsed._replace(query=urllib.parse.urlencode(cleaned_query, doseq=True))
    return strip_control_chars(urllib.parse.urlunparse(parsed))


def is_same_scope(root_url: str, candidate_url: str) -> bool:
    root_host = scope_host(urllib.parse.urlparse(root_url).hostname or "")
    candidate_host = scope_host(urllib.parse.urlparse(candidate_url).hostname or "")
    return bool(root_host and candidate_host and root_host == candidate_host)


def is_yandex_disk_public_url(url: str) -> bool:
    hostname = (urllib.parse.urlparse(url).hostname or "").lower()
    return hostname in {"disk.yandex.ru", "yadi.sk"}


def is_roseltorg_url(url: str) -> bool:
    hostname = (urllib.parse.urlparse(url).hostname or "").lower()
    return hostname.endswith("roseltorg.ru")


def is_roseltorg_procedure_path(path: str) -> bool:
    return bool(re.fullmatch(r"/procedure/[^/?#]+/?", path.lower()))


def is_known_unhelpful_url(url: str) -> bool:
    hostname = (urllib.parse.urlparse(url).hostname or "").lower()
    lowered = url.lower()
    if hostname in {
        "vk.com",
        "hh.ru",
        "t.me",
        "apps.apple.com",
        "play.google.com",
        "appgallery.huawei.com",
        "rustore.ru",
        "www.rustore.ru",
        "2b.pik.ru",
    }:
        return True
    return any(
        token in lowered
        for token in (
            "/ecp",
            "/bg",
            "/education",
            "/knowledge_db",
            "/realty-support",
            "passport.yandex",
            "mc.yandex",
            "/metrika",
            "/player-api",
            "/watch.js",
            "opensearch-",
            "subscribe-options",
            "product-recommendations",
            "tender.pik.ru/feedback",
        )
    )


def looks_like_document_url(url: str) -> bool:
    parsed = urllib.parse.urlparse(url)
    suffix = Path(parsed.path).suffix.lower()
    if suffix in DOCUMENT_EXTENSIONS:
        return True
    lowered = url.lower()
    return any(token in lowered for token in ("/download", "attachment", "filename=", "file=", "/export"))


def looks_like_asset_url(url: str) -> bool:
    parsed = urllib.parse.urlparse(url)
    suffix = Path(parsed.path).suffix.lower()
    if suffix in IGNORED_ASSET_EXTENSIONS:
        return True
    return is_known_unhelpful_url(url)


def is_useful_material_link(url: str) -> bool:
    if not url or looks_like_asset_url(url):
        return False
    parsed = urllib.parse.urlparse(url)
    path = (parsed.path or "/").lower()
    if is_roseltorg_url(url):
        if looks_like_document_url(url):
            return True
        return any(
            token in path
            for token in (
                "/documents",
                "/files",
                "/details",
                "/download",
            )
        ) or is_roseltorg_procedure_path(path)
    if "tender.pik.ru" in (parsed.hostname or "").lower():
        if looks_like_document_url(url):
            return True
        return "/tenders/" in path
    return True


def looks_like_browseable_page(url: str) -> bool:
    if is_known_unhelpful_url(url):
        return False
    parsed = urllib.parse.urlparse(url)
    suffix = Path(parsed.path).suffix.lower()
    if suffix:
        return suffix in {".html", ".htm"}
    path = parsed.path or "/"
    if is_roseltorg_url(url):
        return any(
            token in path.lower()
            for token in ("/documents", "/files", "/details")
        ) or is_roseltorg_procedure_path(path)
    return any(
        token in path.lower()
        for token in ("/tenders/", "/disk/", "/d/", "/folder/", "/file/", "/share/", "/project", "/feedback")
    ) or path.endswith("/")


def should_follow_nested_link(root_url: str, candidate_url: str) -> bool:
    if not is_useful_material_link(candidate_url):
        return False
    if looks_like_document_url(candidate_url):
        return True
    return is_same_scope(root_url, candidate_url) and looks_like_browseable_page(candidate_url)


def make_unique_path(target: Path) -> Path:
    stem, suffix = target.stem, target.suffix
    index = 1
    unique_target = target
    while unique_target.exists():
        unique_target = target.with_name(f"{stem}_{index}{suffix}")
        index += 1
    return unique_target


def is_html_content_type(content_type: str) -> bool:
    return "text/html" in content_type.lower()


def resolve_download_target(output_dir: Path, final_url: str, content_type: str) -> Path:
    ensure_dir(output_dir)
    parsed = urllib.parse.urlparse(final_url)
    candidate_name = urllib.parse.unquote(Path(parsed.path).name or "downloaded")
    candidate_name = slugify(candidate_name, "downloaded")
    target = output_dir / candidate_name
    if not target.suffix:
        if is_html_content_type(content_type):
            target = target.with_suffix(".html")
        else:
            ext = mimetypes.guess_extension(content_type.split(";")[0].strip())
            if ext:
                target = target.with_suffix(ext)
    return make_unique_path(target)


def download_binary_to_path(url: str, target: Path, timeout: int, user_agent: str) -> Optional[DownloadedResource]:
    headers = {"User-Agent": user_agent}
    try:
        if requests is not None:
            response = requests.get(url, timeout=timeout, headers=headers, allow_redirects=True)
            response.raise_for_status()
            content = response.content
            final_url = response.url
            content_type = response.headers.get("Content-Type", "")
            text_content = response.text if is_html_content_type(content_type) else ""
        else:
            request = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(request, timeout=timeout) as response:
                content = response.read()
                final_url = response.geturl()
                content_type = response.headers.get_content_type()
                charset = response.headers.get_content_charset() or "utf-8"
                text_content = content.decode(charset, errors="ignore") if is_html_content_type(content_type) else ""
        ensure_dir(target.parent)
        final_target = make_unique_path(target)
        final_target.write_bytes(content)
        return DownloadedResource(
            source_url=url,
            final_url=final_url,
            local_path=str(final_target),
            content_type=content_type,
            text_content=text_content,
        )
    except (requests.RequestException if requests is not None else urllib.error.URLError, OSError, ValueError):
        return None


def message_has_attachments(message: Message) -> bool:
    for part in message.walk():
        if part.get_filename():
            return True
    return False


def save_message_artifacts(message: Message, raw_message: bytes, base_dir: Path) -> Tuple[List[str], str]:
    ensure_dir(base_dir)
    eml_path = base_dir / "message.eml"
    eml_path.write_bytes(raw_message)
    text_body, html_body = extract_message_bodies(message)
    body_text = text_body.strip() or html_to_text(html_body)
    body_path = base_dir / "message.txt"
    body_path.write_text(body_text, encoding="utf-8")
    attachments = save_attachments(message, base_dir / "attachments")
    return attachments, body_text


def save_attachments(message: Message, attachment_dir: Path) -> List[str]:
    saved: List[str] = []
    for part in message.walk():
        filename = part.get_filename()
        if not filename:
            continue
        payload = part.get_payload(decode=True)
        if payload is None:
            continue
        ensure_dir(attachment_dir)
        safe_name = slugify(decode_mime(filename), "attachment")
        target = attachment_dir / safe_name
        stem, suffix = target.stem, target.suffix
        index = 1
        while target.exists():
            target = attachment_dir / f"{stem}_{index}{suffix}"
            index += 1
        target.write_bytes(payload)
        saved.append(str(target))
    return saved


def download_url(url: str, output_dir: Path, timeout: int, user_agent: str) -> Optional[DownloadedResource]:
    headers = {"User-Agent": user_agent}
    try:
        if requests is not None:
            response = requests.get(url, timeout=timeout, headers=headers, allow_redirects=True)
            response.raise_for_status()
            content = response.content
            final_url = response.url
            content_type = response.headers.get("Content-Type", "")
            target = resolve_download_target(output_dir, final_url, content_type)
            text_content = response.text if is_html_content_type(content_type) else ""
        else:
            request = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(request, timeout=timeout) as response:
                content = response.read()
                final_url = response.geturl()
                content_type = response.headers.get_content_type()
                target = resolve_download_target(output_dir, final_url, content_type)
                charset = response.headers.get_content_charset() or "utf-8"
                text_content = content.decode(charset, errors="ignore") if is_html_content_type(content_type) else ""
        target.write_bytes(content)
        return DownloadedResource(
            source_url=url,
            final_url=final_url,
            local_path=str(target),
            content_type=content_type,
            text_content=text_content,
        )
    except (requests.RequestException if requests is not None else urllib.error.URLError, OSError, ValueError):
        return None


def fetch_json(url: str, params: Dict[str, str], timeout: int, user_agent: str) -> Optional[Dict[str, object]]:
    headers = {"User-Agent": user_agent}
    try:
        if requests is not None:
            response = requests.get(url, params=params, timeout=timeout, headers=headers, allow_redirects=True)
            response.raise_for_status()
            return response.json()
        full_url = url + ("?" + urllib.parse.urlencode(params) if params else "")
        request = urllib.request.Request(full_url, headers=headers)
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode(response.headers.get_content_charset() or "utf-8", errors="ignore"))
    except Exception:
        return None


def download_yandex_disk_public_resource(
    public_url: str,
    output_dir: Path,
    timeout: int,
    user_agent: str,
    max_files: int,
) -> List[DownloadedResource]:
    if max_files <= 0:
        return []

    root_dir = output_dir / slugify(Path(urllib.parse.urlparse(public_url).path).name or "yandex_disk", "yandex_disk")
    api_url = "https://cloud-api.yandex.net/v1/disk/public/resources"
    download_api_url = "https://cloud-api.yandex.net/v1/disk/public/resources/download"
    downloaded: List[DownloadedResource] = []
    visited_paths: set[str] = set()

    def walk(path: str = "") -> None:
        if len(downloaded) >= max_files:
            return
        key = path or "/"
        if key in visited_paths:
            return
        visited_paths.add(key)

        params = {"public_key": public_url, "limit": "1000"}
        if path:
            params["path"] = path
        payload = fetch_json(api_url, params, timeout, user_agent)
        if not payload:
            return

        resource_type = str(payload.get("type", ""))
        resource_name = str(payload.get("name", "")) or "resource"
        if resource_type == "file":
            download_params = {"public_key": public_url}
            if path:
                download_params["path"] = path
            download_payload = fetch_json(download_api_url, download_params, timeout, user_agent)
            href = str((download_payload or {}).get("href", ""))
            if href:
                target = root_dir / slugify(resource_name, "file")
                resource = download_binary_to_path(href, target, timeout, user_agent)
                if resource:
                    downloaded.append(resource)
            return

        items = (((payload.get("_embedded") or {}).get("items")) or [])
        for item in items:
            if len(downloaded) >= max_files:
                break
            item_type = str(item.get("type", ""))
            item_name = str(item.get("name", "")) or "resource"
            item_path = str(item.get("path", ""))
            if item_type == "dir":
                walk(item_path)
                continue
            if item_type != "file":
                continue
            download_params = {"public_key": public_url, "path": item_path}
            download_payload = fetch_json(download_api_url, download_params, timeout, user_agent)
            href = str((download_payload or {}).get("href", ""))
            if not href:
                continue
            relative_parts = [slugify(part, "folder") for part in item_path.replace("disk:/", "").split("/")[:-1] if part]
            target_dir = root_dir
            for part in relative_parts:
                target_dir /= part
            target = target_dir / slugify(item_name, "file")
            resource = download_binary_to_path(href, target, timeout, user_agent)
            if resource:
                downloaded.append(resource)

    walk()
    return downloaded


def download_link_graph(
    root_url: str,
    output_dir: Path,
    timeout: int,
    user_agent: str,
    max_depth: int = 2,
    max_files: int = 20,
) -> List[DownloadedResource]:
    if is_yandex_disk_public_url(root_url):
        return download_yandex_disk_public_resource(root_url, output_dir, timeout, user_agent, max_files)

    downloaded: List[DownloadedResource] = []
    queue: List[Tuple[str, int]] = [(root_url, 0)]
    visited: set[str] = set()
    root_folder = output_dir / slugify(Path(urllib.parse.urlparse(root_url).path).name or urllib.parse.urlparse(root_url).netloc, "link")

    while queue and len(downloaded) < max_files:
        current_url, depth = queue.pop(0)
        normalized_url = normalize_candidate_url(urllib.parse.urldefrag(current_url)[0])
        if normalized_url in visited:
            continue
        visited.add(normalized_url)

        resource = download_url(normalized_url, root_folder, timeout, user_agent)
        if not resource:
            continue
        downloaded.append(resource)

        if depth >= max_depth or not is_html_content_type(resource.content_type):
            continue

        nested_links = extract_links(resource.text_content, resource.text_content, base_url=resource.final_url)
        for nested_link in nested_links:
            clean_link = normalize_candidate_url(urllib.parse.urldefrag(nested_link)[0])
            if clean_link in visited:
                continue
            if not should_follow_nested_link(root_url, clean_link):
                continue
            queue.append((clean_link, depth + 1))

    return downloaded


def extract_text_from_path(path: Path) -> Tuple[str, List[str]]:
    notes: List[str] = []
    suffix = path.suffix.lower()
    if suffix in {".txt", ".csv", ".json", ".xml"}:
        try:
            return path.read_text(encoding="utf-8", errors="ignore"), notes
        except Exception as exc:
            notes.append(f"Не удалось прочитать текстовый файл {path.name}: {exc}")
            return "", notes
    if suffix in {".htm", ".html"}:
        try:
            return html_to_text(path.read_text(encoding="utf-8", errors="ignore")), notes
        except Exception as exc:
            notes.append(f"Не удалось прочитать HTML {path.name}: {exc}")
            return "", notes
    if suffix == ".docx":
        return extract_text_from_docx(path, notes)
    if suffix == ".xlsx":
        return extract_text_from_xlsx(path, notes)
    if suffix == ".pdf":
        return extract_text_from_pdf(path, notes)
    if suffix == ".xls":
        notes.append(f"Файл {path.name} сохранен, но текстовый разбор .xls не реализован без внешних библиотек")
        return "", notes
    return "", notes


def extract_text_from_docx(path: Path, notes: List[str]) -> Tuple[str, List[str]]:
    try:
        with zipfile.ZipFile(path) as archive:
            xml_data = archive.read("word/document.xml")
        root = ET.fromstring(xml_data)
        text_nodes = [node.text for node in root.iter() if node.tag.endswith("}t") and node.text]
        return "\n".join(text_nodes), notes
    except Exception as exc:
        notes.append(f"Не удалось извлечь текст из DOCX {path.name}: {exc}")
        return "", notes


def extract_text_from_xlsx(path: Path, notes: List[str]) -> Tuple[str, List[str]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        chunks: List[str] = []
        for sheet in workbook.worksheets:
            chunks.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell not in (None, "")]
                if values:
                    chunks.append(" | ".join(values))
        return "\n".join(chunks), notes
    except Exception as exc:
        notes.append(f"Не удалось извлечь текст из XLSX {path.name}: {exc}")
        return "", notes


def extract_text_from_pdf(path: Path, notes: List[str]) -> Tuple[str, List[str]]:
    if PDF_READER is None:
        notes.append(f"Файл {path.name} сохранен, но модуль чтения PDF не установлен")
        return "", notes
    try:
        reader = PDF_READER(str(path))
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        return "\n".join(pages), notes
    except Exception as exc:
        notes.append(f"Не удалось извлечь текст из PDF {path.name}: {exc}")
        return "", notes


def classify_sections(texts: Iterable[str]) -> List[str]:
    haystack = " ".join(texts).lower()
    found: List[str] = []
    for section, keywords in SECTION_KEYWORDS.items():
        if any(keyword in haystack for keyword in keywords):
            found.append(section)
    return found or ["прочие разделы"]


def extract_contact_info(text: str, sender: str) -> str:
    phones = re.findall(r"(?:\+7|8)[\s(.-]*\d[\d\s().-]{8,}\d", text)
    emails = re.findall(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, flags=re.IGNORECASE)
    contacts = []
    if sender:
        contacts.append(sender)
    contacts.extend(dict.fromkeys(phones))
    contacts.extend([mail for mail in dict.fromkeys(emails) if mail.lower() not in sender.lower()])
    return "; ".join(contacts[:5])


def extract_first_match(patterns: Sequence[str], text: str) -> str:
    for pattern in patterns:
        match = re.search(pattern, text, flags=re.IGNORECASE)
        if match:
            return match.group(1).strip()
    return ""


def extract_object_name(subject: str, text: str) -> str:
    subject = subject.strip()
    for pattern in (
        r"(?:объект|проект)\s*[:\-]\s*(.+)",
        r"тз\s+по\s+(.+)",
        r"кп\s+по\s+(.+)",
    ):
        match = re.search(pattern, f"{subject}\n{text}", flags=re.IGNORECASE)
        if match:
            return match.group(1).strip()
    return subject[:200]


def summarize_docs(paths: Iterable[str]) -> str:
    names = [Path(path).name for path in paths]
    return "; ".join(names)


def extract_field_bundle(subject: str, body_text: str, combined_text: str, sender: str) -> Dict[str, str]:
    extracted = {}
    extracted["object"] = extract_object_name(subject, combined_text)
    extracted["customer"] = extract_first_match(FIELD_PATTERNS["customer"], combined_text) or sender
    extracted["contacts"] = extract_contact_info(combined_text, sender)
    extracted["address"] = extract_first_match(FIELD_PATTERNS["address"], combined_text)
    extracted["work_type"] = extract_first_match(FIELD_PATTERNS["work_type"], combined_text) or subject
    extracted["stage"] = extract_first_match(FIELD_PATTERNS["stage"], combined_text)
    extracted["deadline"] = extract_first_match(FIELD_PATTERNS["deadline"], combined_text)
    extracted["requirements"] = extract_requirements(combined_text)
    extracted["constraints"] = extract_constraints(combined_text)
    extracted["volumes"] = extract_volumes(combined_text)
    extracted["tech_params"] = extract_tech_params(combined_text)
    extracted["missing_info"] = infer_missing_info(extracted, body_text)
    return extracted


RELEVANCE_STRONG_TOKENS = (
    "проектир",
    "проектно",
    "рабочей документац",
    "разработка пд",
    "разработка рд",
    "техническое задание",
    "исходные данные",
    "коммерческое предложение",
    "смет",
    "тендер",
    "торг",
    "закупк",
    "предагр",
    "архитектур",
    "конструктив",
    "инженерные сети",
    "РїСЂРѕРµРєС‚РёСЂ",
    "РїСЂРѕРµРєС‚РЅРѕ",
    "СЂР°Р±РѕС‡РµР№ РґРѕРєСѓРјРµРЅС‚Р°С†",
    "СЂР°Р·СЂР°Р±РѕС‚РєР° РїРґ",
    "СЂР°Р·СЂР°Р±РѕС‚РєР° СЂРґ",
    "С‚РµС…РЅРёС‡РµСЃРєРѕРµ Р·Р°РґР°РЅРёРµ",
    "РёСЃС…РѕРґРЅС‹Рµ РґР°РЅРЅС‹Рµ",
    "РёСЃС…РѕРґРЅРѕ-",
    "РєРѕРјРјРµСЂС‡РµСЃРєРѕРµ РїСЂРµРґР»РѕР¶РµРЅРёРµ",
    "РєРї",
    "СЃРјРµС‚",
    "С‚РµРЅРґРµСЂ",
    "С‚РѕСЂРі",
    "Р·Р°РєСѓРїРє",
    "РїСЂРµРґР°РіСЂ",
    "Р°СЂС…РёС‚РµРєС‚СѓСЂ",
    "РєРѕРЅСЃС‚СЂСѓРєС‚РёРІ",
    "РёРЅР¶РµРЅРµСЂРЅС‹Рµ СЃРµС‚Рё",
)

RELEVANCE_SUBJECT_TOKENS = (
    "приглашение на участие",
    "принять участие",
    "предагр",
    "РїСЂРёРіР»Р°С€РµРЅРёРµ РЅР° СѓС‡Р°СЃС‚РёРµ",
    "РїСЂРёРЅСЏС‚СЊ СѓС‡Р°СЃС‚РёРµ",
    "fwd__РїСЂРёРіР»Р°С€РµРЅРёРµ",
    "РїСЂРµРґР°РіСЂ",
)

RELEVANCE_LINK_TOKENS = (
    "partner.samolet.ru/tenders",
    "tender.pik.ru/tenders",
    "roseltorg.ru/procedure/",
)

RELEVANCE_FILE_TOKENS = (
    "тз",
    "задани",
    "проект",
    "предагр",
    "исход",
    "смет",
    "кп",
    "С‚Р·",
    "Р·Р°РґР°РЅРё",
    "РїСЂРѕРµРєС‚",
    "РїСЂРµРґР°РіСЂ",
    "РёСЃС…РѕРґ",
    "СЃРјРµС‚",
    "РєРї",
    "Р°СЂ",
    "РєСЂ",
)


def assess_design_request_relevance(
    subject: str,
    sender: str,
    body_text: str,
    combined_text: str,
    attachments: Sequence[str],
    links: Sequence[str],
    downloaded_files: Sequence[str],
) -> Tuple[bool, str]:
    haystack = "\n".join([subject, sender, body_text, combined_text]).lower()
    evidence: List[str] = []

    matched_subject = [token for token in RELEVANCE_SUBJECT_TOKENS if token in subject.lower()]
    if matched_subject:
        evidence.append("СЃРёРіРЅР°Р» РІ С‚РµРјРµ РїРёСЃСЊРјР°")

    matched_text = [token for token in RELEVANCE_STRONG_TOKENS if token in haystack]
    if len(matched_text) >= 2:
        evidence.append("РЅР°Р№РґРµРЅС‹ РїСЂРѕРµРєС‚РЅС‹Рµ/С‚РµРЅРґРµСЂРЅС‹Рµ РєР»СЋС‡РµРІС‹Рµ СЃР»РѕРІР°")

    material_names = " ".join(Path(path).name.lower() for path in [*attachments, *downloaded_files])
    if any(token in material_names for token in RELEVANCE_FILE_TOKENS):
        evidence.append("РЅР°Р№РґРµРЅС‹ РїСЂРѕРµРєС‚РЅС‹Рµ РёРјРµРЅР° С„Р°Р№Р»РѕРІ")

    link_text = " ".join(link.lower() for link in links)
    if any(token in link_text for token in RELEVANCE_LINK_TOKENS):
        evidence.append("СЃСЃС‹Р»РєР° РЅР° С‚РµРЅРґРµСЂРЅСѓСЋ/Р·Р°РєСѓРїРѕС‡РЅСѓСЋ РїР»РѕС‰Р°РґРєСѓ")

    return bool(evidence), "; ".join(dict.fromkeys(evidence)) or "РЅРµС‚ СЏРІРЅРѕРіРѕ Р·Р°РїСЂРѕСЃР° РЅР° РїСЂРѕРµРєС‚РёСЂРѕРІР°РЅРёРµ/РљРџ"


def extract_requirements(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    matched = [line for line in lines if any(token in line.lower() for token in ("треб", "необходимо", "обязательно", "просим"))]
    return " | ".join(matched[:5])


def extract_constraints(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    matched = [line for line in lines if any(token in line.lower() for token in ("огранич", "услови", "лимит", "без", "запрещ"))]
    return " | ".join(matched[:5])


def extract_volumes(text: str) -> str:
    matched = re.findall(r"\b\d+[.,]?\d*\s?(?:м2|м3|км|п.м|шт|га|кв\.?\s?м|тонн?)\b", text, flags=re.IGNORECASE)
    return "; ".join(dict.fromkeys(matched))


def extract_tech_params(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    matched = [line for line in lines if any(token in line.lower() for token in ("мощн", "производительн", "диаметр", "давлен", "расход", "напряжен", "температур"))]
    return " | ".join(matched[:5])


def infer_missing_info(extracted: Dict[str, str], body_text: str) -> str:
    missing = []
    for key, label in (
        ("object", "не указан объект"),
        ("address", "нет адреса или местоположения"),
        ("stage", "нет стадии проектирования"),
        ("deadline", "нет сроков"),
        ("requirements", "нет явных требований"),
    ):
        if not extracted.get(key):
            missing.append(label)
    if "смет" in body_text.lower() or "кп" in body_text.lower():
        for key, label in (
            ("volumes", "нет объемов работ"),
            ("tech_params", "нет ключевых технических параметров"),
        ):
            if not extracted.get(key):
                missing.append(label)
    return "; ".join(missing)


def assess_sufficiency(extracted: Dict[str, str], attachments: Sequence[str], links: Sequence[str], sections: Sequence[str]) -> Tuple[str, str, str]:
    score = 0
    reasons = []
    for key in ("object", "address", "work_type", "stage", "deadline", "requirements"):
        if extracted.get(key):
            score += 1
    if attachments or links:
        score += 1
    if sections and sections != ["прочие разделы"]:
        score += 1
    if not extracted.get("volumes"):
        reasons.append("нет подтвержденных объемов")
    if not extracted.get("tech_params"):
        reasons.append("нет технических параметров")
    if not attachments and not links:
        reasons.append("нет приложенных материалов")
    if score >= 7:
        status = "достаточно"
        risk = "; ".join(reasons[:2]) or "критичных пробелов не выявлено"
        action = "Готовить смету/КП и при необходимости уточнить отдельные параметры"
    elif score >= 4:
        status = "частично достаточно"
        risk = "; ".join(reasons[:3]) or "требуется уточнение части исходных данных"
        action = "Запросить недостающие данные до подготовки точного КП"
    else:
        status = "недостаточно"
        risk = "; ".join(reasons[:3]) or "недостаточно исходных данных"
        action = "Сформировать перечень недостающих исходных данных и запросить их"
    return status, risk, action


def build_report_row(processed_at: str, item: ProcessedMessage) -> List[str]:
    fields = item.extracted_fields
    status, risk, action = assess_sufficiency(fields, item.attachments, item.links, item.sections)
    return [
        processed_at,
        item.mailbox,
        item.uid,
        item.subject,
        fields.get("object", ""),
        fields.get("customer", ""),
        fields.get("contacts", ""),
        fields.get("address", ""),
        fields.get("work_type", ""),
        fields.get("stage", ""),
        fields.get("deadline", ""),
        summarize_docs(item.attachments + item.downloaded_files),
        "; ".join(item.sections),
        "; ".join(item.attachments + item.downloaded_files),
        "; ".join(item.links),
        fields.get("tech_params", ""),
        fields.get("volumes", ""),
        fields.get("requirements", ""),
        fields.get("constraints", ""),
        fields.get("missing_info", ""),
        risk,
        action,
        status,
        item.output_dir,
    ]


def append_rows_to_report(report_path: Path, rows: List[List[str]]) -> None:
    ensure_dir(report_path.parent)
    if report_path.exists():
        workbook = load_workbook(report_path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Исходные данные"
        sheet.append(REPORT_HEADERS)
    for row in rows:
        sheet.append(row)
    workbook.save(report_path)


def processed_aliases(mailbox_name: str, uid: str, message_id: str) -> List[str]:
    aliases = [f"{mailbox_name}:uid:{uid}"]
    if message_id:
        aliases.insert(0, f"{mailbox_name}:{message_id}")
    return aliases


def is_message_processed(state: Dict[str, Dict[str, object]], mailbox_name: str, uid: str, message_id: str) -> bool:
    processed = state.get("processed", {})
    return any(alias in processed for alias in processed_aliases(mailbox_name, uid, message_id))


def mark_message_processed(
    state: Dict[str, Dict[str, object]],
    mailbox_name: str,
    uid: str,
    message_id: str,
    payload: Dict[str, object],
) -> None:
    for alias in processed_aliases(mailbox_name, uid, message_id):
        state["processed"][alias] = payload


def sanitize_sender(sender_header: str) -> str:
    name, address = email.utils.parseaddr(sender_header)
    raw = name or address or sender_header
    return decode_mime(raw)


def process_mailbox(cfg: MailboxConfig, config: Dict[str, object], state: Dict[str, Dict[str, object]]) -> List[ProcessedMessage]:
    print(f"[INFO] Checking mailbox {cfg.name}...", flush=True)
    client = connect_mailbox(cfg)
    results: List[ProcessedMessage] = []
    try:
        uids = search_all_uids(client)
        for uid in uids:
            if is_message_processed(state, cfg.name, uid, ""):
                continue
            raw_message = fetch_message_bytes(client, uid)
            message = email.message_from_bytes(raw_message)
            subject = decode_mime(message.get("Subject"))
            sender = sanitize_sender(message.get("From", ""))
            message_id = (message.get("Message-ID") or "").strip("<>")
            if is_message_processed(state, cfg.name, uid, message_id):
                continue
            body_text, html_body = extract_message_bodies(message)
            links = extract_links(body_text, html_body)
            if not message_has_attachments(message) and not links:
                continue
            date_folder = parse_email_date(message).isoformat()
            sender_slug = slugify(sender, "unknown_sender")
            msg_slug = slugify(subject, hashlib.sha1(f"{cfg.name}:{uid}:{message_id}".encode("utf-8")).hexdigest()[:10])[:80]
            base_dir = Path(str(config["output_root"])) / date_folder / slugify(cfg.name) / sender_slug / msg_slug
            attachments, saved_body_text = save_message_artifacts(message, raw_message, base_dir)
            downloaded_files: List[str] = []
            for link in links:
                resources = download_link_graph(
                    link,
                    base_dir / "downloads",
                    int(config.get("download_timeout_seconds", 30)),
                    str(config.get("user_agent", "mail-intake-automation/1.0")),
                    int(config.get("link_follow_depth", 2)),
                    int(config.get("max_downloaded_files_per_link", 20)),
                )
                if resources:
                    downloaded_files.extend(resource.local_path for resource in resources)
            extracted_texts = [subject, saved_body_text, html_to_text(html_body)]
            notes: List[str] = []
            for path_string in attachments + downloaded_files:
                text, item_notes = extract_text_from_path(Path(path_string))
                if text:
                    extracted_texts.append(text)
                notes.extend(item_notes)
            if links and not downloaded_files:
                notes.append("Ссылки найдены, но вложенные материалы по ним скачать не удалось")
            sections = classify_sections(extracted_texts)
            combined_text = "\n".join(text for text in extracted_texts if text)
            fields = extract_field_bundle(subject, saved_body_text, combined_text, sender)
            is_relevant, relevance_reason = assess_design_request_relevance(
                subject,
                sender,
                saved_body_text,
                combined_text,
                attachments,
                links,
                downloaded_files,
            )
            fields["is_relevant"] = "yes" if is_relevant else "no"
            fields["relevance_reason"] = relevance_reason
            item = ProcessedMessage(
                mailbox=cfg.name,
                uid=uid,
                message_id=message_id,
                subject=subject,
                sender=sender,
                date_folder=date_folder,
                output_dir=str(base_dir.resolve()),
                attachments=[str(Path(path).resolve()) for path in attachments],
                links=links,
                downloaded_files=[str(Path(path).resolve()) for path in downloaded_files],
                sections=sections,
                extracted_fields=fields,
                notes=notes,
            )
            summary_path = base_dir / "analysis.json"
            summary_payload = {
                "mailbox": item.mailbox,
                "uid": item.uid,
                "message_id": item.message_id,
                "subject": item.subject,
                "sender": item.sender,
                "sections": item.sections,
                "fields": item.extracted_fields,
                "attachments": item.attachments,
                "links": item.links,
                "downloaded_files": item.downloaded_files,
                "notes": item.notes,
                "is_relevant": is_relevant,
                "relevance_reason": relevance_reason,
            }
            summary_path.write_text(json.dumps(summary_payload, ensure_ascii=False, indent=2), encoding="utf-8")
            mark_message_processed(state, cfg.name, uid, message_id, {
                "processed_at": dt.datetime.now().isoformat(timespec="seconds"),
                "uid": uid,
                "mailbox": cfg.name,
                "subject": subject,
                "path": str(base_dir.resolve()),
            })
            if is_relevant:
                results.append(item)
    finally:
        try:
            client.close()
        except Exception:
            pass
        client.logout()
    return results


def bootstrap_mailbox_state(cfg: MailboxConfig, state: Dict[str, Dict[str, object]]) -> int:
    print(f"[INFO] Bootstrapping mailbox {cfg.name}...", flush=True)
    client = connect_mailbox(cfg)
    bootstrapped = 0
    try:
        for uid in search_all_uids(client):
            if is_message_processed(state, cfg.name, uid, ""):
                continue
            payload = {
                "processed_at": dt.datetime.now().isoformat(timespec="seconds"),
                "uid": uid,
                "mailbox": cfg.name,
                "subject": "__bootstrap_existing__",
                "path": "",
            }
            mark_message_processed(state, cfg.name, uid, "", payload)
            bootstrapped += 1
    finally:
        try:
            client.close()
        except Exception:
            pass
        client.logout()
    return bootstrapped


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="Hourly intake of design-source materials from IMAP mailboxes")
    parser.add_argument("--config", default="config.json", help="Path to JSON config")
    args = parser.parse_args(argv)

    config_path = Path(args.config)
    if not config_path.exists():
        print(f"Config file not found: {config_path}", file=sys.stderr)
        return 2
    config = load_config(config_path)
    global RUNTIME_CONFIG
    RUNTIME_CONFIG = config
    state_path = Path(str(config.get("state_dir", ".state"))) / "processed_messages.json"
    state = load_state(state_path)
    processed_at = dt.datetime.now().isoformat(timespec="seconds")
    report_date = dt.date.today().isoformat()
    all_items: List[ProcessedMessage] = []
    bootstrap_skip_existing = bool(config.get("bootstrap_skip_existing", True))

    if not state.get("processed") and bootstrap_skip_existing:
        total_bootstrapped = 0
        for mailbox in config["mailboxes"]:
            try:
                total_bootstrapped += bootstrap_mailbox_state(mailbox, state)
                print(f"[INFO] Bootstrapped mailbox {mailbox.name}", flush=True)
            except (socket.timeout, TimeoutError) as exc:
                print(f"[WARN] Mailbox {mailbox.name}: timeout during bootstrap: {exc}", file=sys.stderr, flush=True)
            except Exception as exc:
                print(f"[WARN] Mailbox {mailbox.name}: bootstrap failed: {exc}", file=sys.stderr, flush=True)
        save_state(state_path, state)
        print(json.dumps({
            "processed_messages": 0,
            "bootstrapped_messages": total_bootstrapped,
            "report_date": report_date,
            "state_path": str(state_path.resolve()),
        }, ensure_ascii=False))
        return 0

    for mailbox in config["mailboxes"]:
        try:
            all_items.extend(process_mailbox(mailbox, config, state))
            print(f"[INFO] Completed mailbox {mailbox.name}", flush=True)
        except (socket.timeout, TimeoutError) as exc:
            print(f"[WARN] Mailbox {mailbox.name}: timeout: {exc}", file=sys.stderr, flush=True)
        except Exception as exc:
            print(f"[WARN] Mailbox {mailbox.name}: {exc}", file=sys.stderr, flush=True)

    if all_items:
        report_path = Path(str(config.get("report_root", "reports"))) / f"report-{report_date}.xlsx"
        rows = [build_report_row(processed_at, item) for item in all_items]
        append_rows_to_report(report_path, rows)
    save_state(state_path, state)
    print(json.dumps({
        "processed_messages": len(all_items),
        "report_date": report_date,
        "state_path": str(state_path.resolve()),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
